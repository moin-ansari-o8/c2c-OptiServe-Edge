"""
Document Loader for RAG pipeline.

Loads knowledge sources from various file formats:
  - .txt  (plain text)
  - .md   (markdown)
  - .pdf  (requires `pypdf`)
  - .docx (requires `python-docx`)

Pipeline:
    Documents on disk  →  loader.py  →  Document objects in memory
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import List

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------

@dataclass
class Document:
    """A loaded document with its raw text content and metadata."""

    content: str
    metadata: dict = field(default_factory=dict)

    @property
    def source(self) -> str:
        return self.metadata.get("source", "unknown")

    @property
    def filename(self) -> str:
        return self.metadata.get("filename", "unknown")

    def __repr__(self) -> str:
        preview = self.content[:60] + "…" if len(self.content) > 60 else self.content
        return f"Document(source={self.source!r}, len={len(self.content)}, preview={preview!r})"


# ---------------------------------------------------------------------------
# Loader
# ---------------------------------------------------------------------------

class DocumentLoader:
    """Load documents from files or directories.

    Supported extensions: .txt, .md, .pdf, .docx
    """

    SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}

    # ── public API ──────────────────────────────────────────────────────

    def load(self, path: str) -> List[Document]:
        """Load document(s) from a file or directory path.

        If *path* is a directory every supported file inside it (recursively)
        is loaded.  Unsupported files are silently skipped when scanning a
        directory; calling with an unsupported *file* raises ``ValueError``.
        """
        p = Path(path)
        if p.is_file():
            return [self._load_file(p)]
        if p.is_dir():
            return self._load_directory(p)
        raise FileNotFoundError(f"Path not found: {path}")

    # ── internals ───────────────────────────────────────────────────────

    def _load_directory(self, directory: Path) -> List[Document]:
        """Recursively load all supported documents from *directory*."""
        documents: List[Document] = []
        for ext in sorted(self.SUPPORTED_EXTENSIONS):
            for file_path in sorted(directory.rglob(f"*{ext}")):
                try:
                    documents.append(self._load_file(file_path))
                    logger.info("Loaded: %s", file_path)
                except Exception as exc:  # noqa: BLE001
                    logger.warning("Skipped %s: %s", file_path, exc)
        return documents

    def _load_file(self, file_path: Path) -> Document:
        """Load a single file, dispatching by extension."""
        ext = file_path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {ext}")

        metadata = {
            "source": str(file_path),
            "filename": file_path.name,
            "extension": ext,
        }

        loader_map = {
            ".txt": self._load_text,
            ".md": self._load_text,
            ".pdf": self._load_pdf,
            ".docx": self._load_docx,
        }

        content = loader_map[ext](file_path)
        return Document(content=content, metadata=metadata)

    # ── format-specific loaders ─────────────────────────────────────────

    @staticmethod
    def _load_text(path: Path) -> str:
        return path.read_text(encoding="utf-8")

    @staticmethod
    def _load_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError(
                "Install 'pypdf' to load PDF files:  pip install pypdf"
            )
        reader = PdfReader(str(path))
        pages = [page.extract_text() for page in reader.pages if page.extract_text()]
        return "\n\n".join(pages)

    @staticmethod
    def _load_docx(path: Path) -> str:
        try:
            import docx  # python-docx
        except ImportError:
            raise ImportError(
                "Install 'python-docx' to load DOCX files:  pip install python-docx"
            )
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
